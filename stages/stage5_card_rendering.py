"""
Stage 5: Card Rendering -- Generate process card as Word .docx

Takes a process route (list of operation dicts from Stage 4) and part info,
renders a formatted A4 landscape Word document with header, process table,
auto-inserted inspection rows, and footer signature blocks.
"""

from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

STAGE_NAME = "Stage5_CardRendering"

# Non-machining step types that should NOT get an inspection row after them
_SKIP_INSPECTION_TYPES = {"领料", "入库", "检验", "包装", "清洗", "去毛刺"}

# Column definitions: (header_text, width_cm)
_COLUMNS = [
    ("工序号", 1.5),
    ("工种名称", 2.0),
    ("工序内容", 14.0),
    ("设备", 3.0),
    ("工艺装备", 3.0),
]


# ---------------------------------------------------------------------------
# Helper: set cell borders via XML
# ---------------------------------------------------------------------------
def _set_cell_border(cell, **kwargs):
    """
    Set cell border properties.

    Usage:
        _set_cell_border(cell,
            top={"sz": 4, "val": "single", "color": "000000"},
            bottom={"sz": 4, "val": "single", "color": "000000"},
            start={"sz": 4, "val": "single", "color": "000000"},
            end={"sz": 4, "val": "single", "color": "000000"},
        )
    """
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()

    tc_borders = OxmlElement("w:tcBorders")

    for edge_name, attrs in kwargs.items():
        edge_el = OxmlElement(f"w:{edge_name}")
        edge_el.set(qn("w:sz"), str(attrs.get("sz", 4)))
        edge_el.set(qn("w:val"), attrs.get("val", "single"))
        edge_el.set(qn("w:color"), attrs.get("color", "000000"))
        edge_el.set(qn("w:space"), "0")
        tc_borders.append(edge_el)

    tc_pr.append(tc_borders)


def _apply_all_borders(cell):
    """Apply thin single borders on all four sides of a cell."""
    border = {"sz": 4, "val": "single", "color": "000000"}
    _set_cell_border(cell, top=border, bottom=border, start=border, end=border)


# ---------------------------------------------------------------------------
# Helper: font styling
# ---------------------------------------------------------------------------
def _set_run_font(run, font_name: str = "SimSun", font_size_pt: float = 10):
    """Set both Western and East Asian font on a run."""
    run.font.name = font_name
    run.font.size = Pt(font_size_pt)
    rpr = run._element.get_or_add_rPr()
    r_fonts = rpr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        rpr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)


def _add_cell_text(cell, text: str, font_size: float = 10,
                   bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT):
    """Write text into a table cell with consistent formatting."""
    cell.text = ""  # clear default paragraph
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run(str(text))
    _set_run_font(run, font_size_pt=font_size)
    run.bold = bold


# ---------------------------------------------------------------------------
# Document layout helpers
# ---------------------------------------------------------------------------
def _set_landscape_a4(document):
    """Set the document to A4 landscape with reasonable margins."""
    section = document.sections[0]
    # A4: 210mm x 297mm  -- landscape swaps width/height
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)


def _add_header_section(document, part_info: dict):
    """Add a header block with part name, material, blank spec, and date."""
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("机械加工工艺过程卡")
    _set_run_font(run, font_size_pt=16)
    run.bold = True

    # Info table: 1 row x 4 columns
    info_table = document.add_table(rows=1, cols=4)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    labels_values = [
        ("零件名称", part_info.get("part_type", "")),
        ("材料", part_info.get("material", "")),
        ("毛坯规格", part_info.get("blank_info", "")),
        ("编制日期", datetime.now().strftime("%Y-%m-%d")),
    ]

    row = info_table.rows[0]
    for idx, (label, value) in enumerate(labels_values):
        cell = row.cells[idx]
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_run = para.add_run(f"{label}: ")
        _set_run_font(label_run, font_size_pt=10)
        label_run.bold = True
        value_run = para.add_run(str(value))
        _set_run_font(value_run, font_size_pt=10)
        _apply_all_borders(cell)

    # Blank line between header and main table
    document.add_paragraph("")


def _needs_inspection(step: dict) -> bool:
    """Determine whether a step should be followed by an inspection row."""
    step_name = step.get("name", "")
    step_type = step.get("type", "")
    combined = f"{step_name}{step_type}"
    for skip in _SKIP_INSPECTION_TYPES:
        if skip in combined:
            return False
    return True


def _build_process_table(document, process_route: list[dict]) -> int:
    """
    Build the main process table and return the number of data rows added.
    """
    num_cols = len(_COLUMNS)

    # Count rows: each step + optional inspection row
    data_rows = 0
    for step in process_route:
        data_rows += 1
        if _needs_inspection(step):
            data_rows += 1

    table = document.add_table(rows=1 + data_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # --- Set column widths ---
    for col_idx, (_, width_cm) in enumerate(_COLUMNS):
        for row in table.rows:
            row.cells[col_idx].width = Cm(width_cm)

    # --- Header row ---
    header_row = table.rows[0]
    for col_idx, (header_text, _) in enumerate(_COLUMNS):
        cell = header_row.cells[col_idx]
        _add_cell_text(cell, header_text, font_size=10, bold=True,
                       align=WD_ALIGN_PARAGRAPH.CENTER)
        _apply_all_borders(cell)

    # --- Data rows ---
    row_idx = 1
    inspection_seq = 0  # running sequence for inspection numbering
    for step in process_route:
        seq_num = step.get("seq", "")
        step_name = step.get("name", "")

        # Join content list with newlines
        content_raw = step.get("content", [])
        if isinstance(content_raw, list):
            content_text = "\n".join(str(item) for item in content_raw)
        else:
            content_text = str(content_raw)

        equipment = step.get("equipment", "")
        tooling = step.get("tooling", "")

        values = [str(seq_num), step_name, content_text, equipment, tooling]
        for col_idx, value in enumerate(values):
            cell = table.rows[row_idx].cells[col_idx]
            align = (WD_ALIGN_PARAGRAPH.CENTER
                     if col_idx in (0, 1) else WD_ALIGN_PARAGRAPH.LEFT)
            _add_cell_text(cell, value, font_size=10, bold=False, align=align)
            _apply_all_borders(cell)
        row_idx += 1

        # --- Auto-inserted inspection row ---
        if _needs_inspection(step):
            inspection_seq += 1
            insp_values = [
                "",
                "检验",
                "按图纸及工艺要求检验本工序内容",
                "",
                "",
            ]
            for col_idx, value in enumerate(insp_values):
                cell = table.rows[row_idx].cells[col_idx]
                align = (WD_ALIGN_PARAGRAPH.CENTER
                         if col_idx in (0, 1) else WD_ALIGN_PARAGRAPH.LEFT)
                _add_cell_text(cell, value, font_size=10, bold=False, align=align)
                _apply_all_borders(cell)
            row_idx += 1

    return data_rows


def _add_footer_section(document):
    """Add signature blocks for drafter, reviewer, and approver."""
    document.add_paragraph("")  # spacer

    footer_table = document.add_table(rows=2, cols=3)
    footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    role_labels = ["编制", "审核", "批准"]
    for col_idx, label in enumerate(role_labels):
        cell = footer_table.rows[0].cells[col_idx]
        _add_cell_text(cell, f"{label}:          ", font_size=10,
                       bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _apply_all_borders(cell)

        date_cell = footer_table.rows[1].cells[col_idx]
        _add_cell_text(date_cell, "日期:          ", font_size=10,
                       bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        _apply_all_borders(date_cell)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------
def render_process_card(
    process_route: list[dict],
    part_info: dict,
    output_path: str,
    tracer: Any,
) -> str:
    """
    Render a process route into a formatted Word .docx process card.

    Args:
        process_route: List of operation dicts from Stage 4. Each dict has
            keys: seq, name, content (list[str]), equipment, tooling, type
        part_info: Dict with keys: part_type, material, blank_info
        output_path: File path for the generated .docx
        tracer: PipelineTracer instance for logging

    Returns:
        Path to the generated .docx file.
    """
    tracer.begin_stage(STAGE_NAME, input_data={
        "num_steps": len(process_route),
        "output_path": output_path,
        "part_type": part_info.get("part_type", ""),
    })

    try:
        # Ensure output directory exists
        out_path = Path(output_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)

        # Create document
        document = Document()
        _set_landscape_a4(document)

        # Header
        _add_header_section(document, part_info)

        # Main process table
        total_rows = _build_process_table(document, process_route)

        # Footer
        _add_footer_section(document)

        # Save
        document.save(str(out_path))

        tracer.log_reasoning(
            f"Process card rendered: {out_path} | "
            f"Total table rows (incl. inspection): {total_rows} | "
            f"Process steps: {len(process_route)}"
        )
        tracer.end_stage(output_data={
            "output_file": str(out_path),
            "total_rows": total_rows,
            "process_steps": len(process_route),
        })

        return str(out_path)

    except Exception as exc:
        error_msg = f"Card rendering failed: {type(exc).__name__}: {exc}"
        tracer.log_error(error_msg)
        tracer.end_stage(output_data={"error": error_msg})
        raise
